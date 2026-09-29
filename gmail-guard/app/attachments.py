"""Text aus Anhängen ziehen. Nichts wird ausgeführt, nur gelesen."""
import io

from common.textutil import html_to_text

TEXT_EXT = (".txt", ".csv", ".md", ".json", ".xml", ".ics", ".log", ".eml", ".vcf")
HTML_EXT = (".html", ".htm")
MAX_PDF_PAGES = 100


def extract_text(data: bytes, filename: str, mime: str):
    """Gibt (text, hinweis) zurück. text=None, wenn das Format nicht unterstützt ist."""
    name = (filename or "").lower()
    mime = (mime or "").lower()

    if mime == "application/pdf" or name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return None, "PDF ist passwortgeschützt."
        pages = reader.pages[:MAX_PDF_PAGES]
        parts, found_text = [], False
        for i, page in enumerate(pages, 1):
            try:
                page_text = page.extract_text() or ""
                found_text = found_text or bool(page_text.strip())
                parts.append(f"--- Seite {i} ---\n{page_text}")
            except Exception:
                parts.append(f"--- Seite {i} --- [nicht lesbar]")
        hint = ""
        if len(reader.pages) > MAX_PDF_PAGES:
            hint = f"Nur die ersten {MAX_PDF_PAGES} von {len(reader.pages)} Seiten gelesen."
        text = "\n".join(parts)
        if not found_text:
            hint = (hint + " Kein Text gefunden – vermutlich ein Scan (Bild-PDF).").strip()
        return text, hint

    if name.endswith(".docx") or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts), ""

    if name.endswith(HTML_EXT) or mime == "text/html":
        return html_to_text(data.decode("utf-8", errors="replace")), ""

    if name.endswith(TEXT_EXT) or mime.startswith("text/"):
        return data.decode("utf-8", errors="replace"), ""

    return None, f"Format nicht unterstützt ({mime or 'unbekannt'}). Unterstützt: PDF, DOCX, HTML, Text/CSV/JSON."
